# plugins/office/plugin.py
"""Office capability plugin implementation for Word, PowerPoint, and Office automation."""

import logging
import os
import re
import subprocess
from typing import Any, Dict, List, Optional

import docx
from docx.shared import Inches
import pptx

from plugins.base import BasePlugin
from plugins.manifest import PluginManifest, PluginResult
from .commands import ACTIONS
from .permissions import PERMISSIONS_REQUIRED

logger = logging.getLogger("jarvis.plugins.office")


class OfficePlugin(BasePlugin):
    """Controls Microsoft Office suite operations (Word, PowerPoint, Office docs)."""

    def get_manifest(self) -> PluginManifest:
        return PluginManifest(
            name="office",
            version="1.0.0",
            description="Microsoft Office suite automation (Word, PowerPoint, documents, metadata, and tables).",
            author="Jarvis Team",
            permissions=PERMISSIONS_REQUIRED,
            actions=ACTIONS
        )

    def execute(self, action_name: str, parameters: Optional[Dict[str, Any]] = None) -> PluginResult:
        parameters = parameters or {}

        if action_name == "create_word_doc":
            title = str(parameters.get("title", "Untitled Document")).strip()
            content = str(parameters.get("content", "")).strip()
            file_path = str(parameters.get("file_path", "")).strip()
            if not file_path:
                safe_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).rstrip()
                file_path = f"{safe_title or 'document'}.docx"

            try:
                os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
                doc = docx.Document()
                doc.add_heading(title, 0)
                if content:
                    doc.add_paragraph(content)
                doc.save(file_path)
            except Exception:
                with open(file_path.replace(".docx", ".txt") if not file_path.endswith(".txt") else file_path, "w", encoding="utf-8") as f:
                    f.write(f"{title}\n{'='*len(title)}\n\n{content}\n")

            return PluginResult(
                success=True,
                data={"file_path": file_path, "title": title, "message": f"Word document '{title}' created at: {file_path}"}
            )

        if action_name == "create_powerpoint":
            title = str(parameters.get("title", "Untitled Presentation")).strip()
            slides_data = parameters.get("slides", [])
            file_path = str(parameters.get("file_path", "")).strip()
            if not file_path:
                safe_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).rstrip()
                file_path = f"{safe_title or 'presentation'}.pptx"

            try:
                os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
                prs = pptx.Presentation()
                title_slide_layout = prs.slide_layouts[0]
                slide = prs.slides.add_slide(title_slide_layout)
                slide.shapes.title.text = title

                if isinstance(slides_data, list):
                    content_layout = prs.slide_layouts[1]
                    for s_item in slides_data:
                        if isinstance(s_item, dict):
                            s_title = s_item.get("title", "Slide")
                            s_bullets = s_item.get("bullets", [])
                            n_slide = prs.slides.add_slide(content_layout)
                            n_slide.shapes.title.text = s_title
                            body_shape = n_slide.shapes.placeholders[1]
                            tf = body_shape.text_frame
                            for idx, b in enumerate(s_bullets):
                                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                                p.text = str(b)
                        elif isinstance(s_item, str):
                            n_slide = prs.slides.add_slide(content_layout)
                            n_slide.shapes.title.text = s_item

                prs.save(file_path)
                return PluginResult(
                    success=True,
                    data={"file_path": file_path, "title": title, "message": f"PowerPoint presentation '{title}' created at: {file_path}"}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to create PowerPoint presentation: {e}")

        if action_name == "open_office_app":
            app = str(parameters.get("app_name", "word")).lower().strip()
            try:
                from core.apps import get_application_registry, get_installed_app_detector
                app_reg = get_application_registry()
                app_def = app_reg.get_app(app) or app_reg.find_by_alias(app)
                detector = get_installed_app_detector()

                if app_def:
                    launch_cmd = detector.get_launch_command(app_def, is_new=parameters.get("is_new", False))
                    if launch_cmd:
                        target_cmd, args, strategy = launch_cmd
                        if strategy == "executable":
                            subprocess.Popen([target_cmd] + args, cwd=os.path.dirname(target_cmd))
                        elif strategy == "shortcut":
                            os.startfile(target_cmd)
                        else:
                            subprocess.Popen(f'start "" "{target_cmd}"', shell=True)
                        return PluginResult(
                            success=True,
                            data={"app": app, "display_name": app_def.display_name, "message": f"'{app_def.display_name}' open kar diya hai, Boss."}
                        )
            except Exception:
                pass

            app_exec = {
                "word": "winword",
                "powerpoint": "powerpnt",
                "ppt": "powerpnt",
                "outlook": "outlook",
                "excel": "excel"
            }.get(app, app)
            try:
                subprocess.Popen(f'start "" "{app_exec}"', shell=True)
                return PluginResult(success=True, data={"app": app, "message": f"'{app.capitalize()}' open kar diya hai, Boss."})
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to open '{app}': {e}")

        if action_name == "add_slide_to_presentation":
            file_path = str(parameters.get("file_path", "")).strip()
            title = str(parameters.get("title", "New Slide")).strip()
            bullets = parameters.get("bullets", [])
            layout_index = int(parameters.get("layout_index", 1))

            if not file_path:
                return PluginResult(success=False, error="file_path is required.")

            try:
                if os.path.exists(file_path):
                    prs = pptx.Presentation(file_path)
                else:
                    prs = pptx.Presentation()

                layout = prs.slide_layouts[min(layout_index, len(prs.slide_layouts) - 1)]
                slide = prs.slides.add_slide(layout)
                if slide.shapes.title:
                    slide.shapes.title.text = title

                if bullets:
                    if len(slide.shapes.placeholders) > 1:
                        tf = slide.shapes.placeholders[1].text_frame
                        bullet_list = bullets if isinstance(bullets, list) else [str(bullets)]
                        for idx, b in enumerate(bullet_list):
                            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                            p.text = str(b)

                prs.save(file_path)
                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "title": title,
                        "total_slides": len(prs.slides),
                        "message": f"Added slide '{title}' to presentation (total: {len(prs.slides)} slides)."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to add slide: {e}")

        if action_name == "get_presentation_summary":
            file_path = str(parameters.get("file_path", "")).strip()
            if not file_path or not os.path.exists(file_path):
                return PluginResult(success=False, error="Valid presentation file_path is required.")

            try:
                prs = pptx.Presentation(file_path)
                slide_titles = []
                total_words = 0
                for idx, slide in enumerate(prs.slides):
                    stitle = slide.shapes.title.text if slide.shapes.title and slide.shapes.title.text else f"Slide {idx+1}"
                    slide_titles.append(stitle)
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                total_words += len(p.text.split())

                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "slide_count": len(prs.slides),
                        "slide_titles": slide_titles,
                        "estimated_word_count": total_words,
                        "message": f"Presentation has {len(prs.slides)} slide(s) and ~{total_words} words."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to summarize presentation: {e}")

        if action_name == "convert_text_to_presentation":
            outline_text = str(parameters.get("outline_text", "")).strip()
            output_path = str(parameters.get("output_path", "")).strip()
            main_title = str(parameters.get("title", "Presentation")).strip()

            if not outline_text or not output_path:
                return PluginResult(success=False, error="outline_text and output_path are required.")

            try:
                os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
                prs = pptx.Presentation()

                # Title slide
                title_slide = prs.slides.add_slide(prs.slide_layouts[0])
                title_slide.shapes.title.text = main_title

                # Parse lines
                content_layout = prs.slide_layouts[1]
                current_slide = None
                current_tf = None

                for line in outline_text.splitlines():
                    line = line.strip()
                    if not line:
                        continue
                    if line.startswith("#") or line.endswith(":"):
                        # New slide header
                        header_text = line.lstrip("#").rstrip(":").strip()
                        current_slide = prs.slides.add_slide(content_layout)
                        current_slide.shapes.title.text = header_text
                        current_tf = current_slide.shapes.placeholders[1].text_frame
                        current_tf.text = ""
                    else:
                        bullet_text = line.lstrip("-*•").strip()
                        if current_slide is None:
                            current_slide = prs.slides.add_slide(content_layout)
                            current_slide.shapes.title.text = "Overview"
                            current_tf = current_slide.shapes.placeholders[1].text_frame
                            current_tf.text = ""
                        p = current_tf.add_paragraph() if current_tf.text else current_tf.paragraphs[0]
                        p.text = bullet_text

                prs.save(output_path)
                return PluginResult(
                    success=True,
                    data={
                        "output_path": output_path,
                        "slide_count": len(prs.slides),
                        "message": f"Created presentation with {len(prs.slides)} slide(s) from outline."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Text to presentation conversion failed: {e}")

        if action_name == "extract_text_from_presentation":
            file_path = str(parameters.get("file_path", "")).strip()
            if not file_path or not os.path.exists(file_path):
                return PluginResult(success=False, error="Valid file_path is required.")

            try:
                prs = pptx.Presentation(file_path)
                extracted_slides = []
                for idx, slide in enumerate(prs.slides):
                    lines = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    lines.append(p.text.strip())
                    extracted_slides.append({
                        "slide_number": idx + 1,
                        "text": "\n".join(lines)
                    })

                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "slides_extracted": len(extracted_slides),
                        "slides": extracted_slides,
                        "message": f"Extracted text from {len(extracted_slides)} slide(s)."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to extract presentation text: {e}")

        if action_name == "add_table_to_document":
            file_path = str(parameters.get("file_path", "")).strip()
            table_data = parameters.get("table_data", [])
            has_header = bool(parameters.get("has_header", True))

            if not file_path or not table_data:
                return PluginResult(success=False, error="file_path and table_data are required.")

            try:
                if os.path.exists(file_path):
                    doc = docx.Document(file_path)
                else:
                    doc = docx.Document()

                rows_cnt = len(table_data)
                cols_cnt = max(len(r) for r in table_data) if rows_cnt > 0 else 0
                table = doc.add_table(rows=rows_cnt, cols=cols_cnt)
                table.style = "Table Grid"

                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = str(val)
                        if r_idx == 0 and has_header:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True

                doc.save(file_path)
                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "rows": rows_cnt,
                        "cols": cols_cnt,
                        "message": f"Added {rows_cnt}x{cols_cnt} table to document '{os.path.basename(file_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to add table to document: {e}")

        if action_name == "insert_image_into_document":
            file_path = str(parameters.get("file_path", "")).strip()
            image_path = str(parameters.get("image_path", "")).strip()
            caption = str(parameters.get("caption", "")).strip()
            width_inches = float(parameters.get("width_inches", 4.0))

            if not file_path or not image_path:
                return PluginResult(success=False, error="file_path and image_path are required.")
            if not os.path.exists(image_path):
                return PluginResult(success=False, error=f"Image file '{image_path}' does not exist.")

            try:
                if os.path.exists(file_path):
                    doc = docx.Document(file_path)
                else:
                    doc = docx.Document()

                doc.add_picture(image_path, width=Inches(width_inches))
                if caption:
                    p = doc.add_paragraph(caption)
                    p.italic = True

                doc.save(file_path)
                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "image_path": image_path,
                        "caption": caption,
                        "message": f"Inserted image into '{os.path.basename(file_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to insert image: {e}")

        if action_name == "get_document_metadata":
            file_path = str(parameters.get("file_path", "")).strip()
            if not file_path or not os.path.exists(file_path):
                return PluginResult(success=False, error="Valid file_path is required.")

            try:
                doc = docx.Document(file_path)
                props = doc.core_properties
                p_count = len(doc.paragraphs)
                t_count = len(doc.tables)
                word_count = sum(len(p.text.split()) for p in doc.paragraphs)

                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "author": props.author or "",
                        "title": props.title or "",
                        "subject": props.subject or "",
                        "comments": props.comments or "",
                        "paragraphs": p_count,
                        "tables": t_count,
                        "word_count": word_count,
                        "message": f"Document '{os.path.basename(file_path)}': {p_count} paragraphs, {word_count} words."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to get metadata: {e}")

        if action_name == "set_document_metadata":
            file_path = str(parameters.get("file_path", "")).strip()
            author = str(parameters.get("author", "")).strip()
            title = str(parameters.get("title", "")).strip()
            comments = str(parameters.get("comments", "")).strip()

            if not file_path or not os.path.exists(file_path):
                return PluginResult(success=False, error="Valid file_path is required.")

            try:
                doc = docx.Document(file_path)
                if author:
                    doc.core_properties.author = author
                if title:
                    doc.core_properties.title = title
                if comments:
                    doc.core_properties.comments = comments
                doc.save(file_path)

                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "author": author,
                        "title": title,
                        "message": f"Updated metadata on '{os.path.basename(file_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to set metadata: {e}")

        if action_name == "merge_word_documents":
            doc_paths = parameters.get("file_paths", [])
            output_path = str(parameters.get("output_path", "")).strip()

            if not doc_paths or not output_path:
                return PluginResult(success=False, error="file_paths list and output_path are required.")

            try:
                master_doc = docx.Document()
                merged_count = 0
                for p in doc_paths:
                    if os.path.exists(p):
                        sub_doc = docx.Document(p)
                        for para in sub_doc.paragraphs:
                            master_doc.add_paragraph(para.text, style=para.style.name if para.style else None)
                        for table in sub_doc.tables:
                            new_table = master_doc.add_table(rows=len(table.rows), cols=len(table.columns))
                            for r_i, row in enumerate(table.rows):
                                for c_i, cell in enumerate(row.cells):
                                    new_table.cell(r_i, c_i).text = cell.text
                        master_doc.add_page_break()
                        merged_count += 1

                master_doc.save(output_path)
                return PluginResult(
                    success=True,
                    data={
                        "output_path": output_path,
                        "merged_count": merged_count,
                        "message": f"Merged {merged_count} documents into '{os.path.basename(output_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Document merge failed: {e}")

        if action_name == "export_document_to_text":
            file_path = str(parameters.get("file_path", "")).strip()
            output_path = str(parameters.get("output_path", "")).strip()

            if not file_path or not os.path.exists(file_path):
                return PluginResult(success=False, error="Valid file_path is required.")

            try:
                doc = docx.Document(file_path)
                lines = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        lines.append(p.text)
                for table in doc.tables:
                    for row in table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))

                full_text = "\n\n".join(lines)
                if output_path:
                    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
                    with open(output_path, "w", encoding="utf-8") as f:
                        f.write(full_text)

                return PluginResult(
                    success=True,
                    data={
                        "file_path": file_path,
                        "output_path": output_path or None,
                        "paragraphs_count": len(lines),
                        "text_length": len(full_text),
                        "preview": full_text[:300],
                        "message": f"Exported {len(lines)} sections from document."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Text export failed: {e}")

        return PluginResult(success=False, error=f"Unknown action '{action_name}' in Office plugin.")

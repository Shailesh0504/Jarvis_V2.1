# plugins/excel/plugin.py
"""Excel Plugin lifecycle and executor using openpyxl and csv."""

import csv
import json
import logging
import os
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from plugins.base import BasePlugin
from plugins.manifest import PluginManifest, PluginResult
from .commands import ACTIONS
from .permissions import PERMISSIONS_REQUIRED

logger = logging.getLogger("jarvis.plugins.excel")


class ExcelPlugin(BasePlugin):
    """Handles Excel workbook operations, formulas, charts, and analysis."""

    def get_manifest(self) -> PluginManifest:
        return PluginManifest(
            name="excel",
            version="1.0.0",
            description="Excel workbook automation, spreadsheet reading, chart generation, and data analysis.",
            author="Jarvis Team",
            permissions=PERMISSIONS_REQUIRED,
            actions=ACTIONS
        )

    def execute(self, action_name: str, parameters: Optional[Dict[str, Any]] = None) -> PluginResult:
        parameters = parameters or {}

        if action_name == "read_sheet":
            file_path = str(parameters.get("path") or parameters.get("filename", "")).strip()
            sheet_name = str(parameters.get("sheet_name", "")).strip()
            if not file_path:
                return PluginResult(success=False, error="File path is required.")
            if not os.path.exists(file_path):
                return PluginResult(success=False, error=f"File '{file_path}' does not exist.")

            try:
                rows_data = []
                headers = []
                if file_path.lower().endswith((".xlsx", ".xlsm", ".xltx")):
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                    sheet_name = ws.title
                    for row in ws.iter_rows(values_only=True):
                        rows_data.append([str(c) if c is not None else "" for c in row])
                    wb.close()
                else:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        reader = csv.reader(f)
                        for row in reader:
                            rows_data.append(row)
                    sheet_name = "CSV"

                if rows_data:
                    headers = rows_data[0]

                return PluginResult(
                    success=True,
                    data={
                        "path": file_path,
                        "sheet": sheet_name,
                        "headers": headers,
                        "rows_read": len(rows_data),
                        "preview": rows_data[:10],
                        "message": f"Read {len(rows_data)} rows from '{os.path.basename(file_path)}' [{sheet_name}]."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to read sheet: {e}")

        if action_name == "create_sheet":
            file_path = str(parameters.get("path") or parameters.get("filename", "")).strip()
            if not file_path:
                return PluginResult(success=False, error="File path is required.")

            try:
                dir_name = os.path.dirname(os.path.abspath(file_path))
                if dir_name:
                    os.makedirs(dir_name, exist_ok=True)

                if file_path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.Workbook()
                    wb.save(file_path)
                else:
                    with open(file_path, "w", newline="", encoding="utf-8") as f:
                        pass

                return PluginResult(
                    success=True,
                    data={"path": file_path, "message": f"Created new spreadsheet at '{file_path}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to create spreadsheet: {e}")

        if action_name == "write_sheet":
            file_path = str(parameters.get("path") or parameters.get("filename", "")).strip()
            rows_data = parameters.get("data", [])
            sheet_name = str(parameters.get("sheet_name", "")).strip()
            if not file_path:
                return PluginResult(success=False, error="File path is required.")

            try:
                dir_name = os.path.dirname(os.path.abspath(file_path))
                if dir_name:
                    os.makedirs(dir_name, exist_ok=True)

                if file_path.lower().endswith((".xlsx", ".xlsm")):
                    if os.path.exists(file_path):
                        wb = openpyxl.load_workbook(file_path)
                    else:
                        wb = openpyxl.Workbook()
                    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                    if isinstance(rows_data, list):
                        for row in rows_data:
                            if isinstance(row, list):
                                ws.append(row)
                            else:
                                ws.append([row])
                    wb.save(file_path)
                else:
                    with open(file_path, "a", newline="", encoding="utf-8") as f:
                        writer = csv.writer(f)
                        if isinstance(rows_data, list):
                            for row in rows_data:
                                if isinstance(row, list):
                                    writer.writerow(row)
                                else:
                                    writer.writerow([str(row)])
                        else:
                            writer.writerow([str(rows_data)])

                count = len(rows_data) if isinstance(rows_data, list) else 1
                return PluginResult(
                    success=True,
                    data={
                        "path": file_path,
                        "rows_written": count,
                        "message": f"Successfully wrote {count} row(s) to '{os.path.basename(file_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to write to sheet: {e}")

        if action_name == "calculate_column":
            file_path = str(parameters.get("path", "")).strip()
            col_target = str(parameters.get("column_name", "sales")).strip().lower()
            operation = str(parameters.get("operation", "sum")).strip().lower()
            sheet_name = str(parameters.get("sheet_name", "")).strip()

            if not file_path:
                return PluginResult(success=False, error="File path parameter is required.")
            if not os.path.exists(file_path):
                return PluginResult(success=False, error=f"File '{file_path}' does not exist.")

            try:
                values = []
                matched_col_name = None

                if file_path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                    headers = [str(c.value or "").strip() for c in ws[1]]
                    col_idx = -1
                    for idx, h in enumerate(headers):
                        if col_target in h.lower() or h.lower() in col_target:
                            col_idx = idx + 1
                            matched_col_name = h
                            break
                    if col_idx == -1:
                        col_idx = len(headers)
                        matched_col_name = headers[-1] if headers else "Column"

                    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
                        val = row[0]
                        if val is not None:
                            try:
                                values.append(float(str(val).replace("$", "").replace("₹", "").replace(",", "").strip()))
                            except ValueError:
                                pass
                    wb.close()
                else:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        reader = csv.reader(f)
                        headers = None
                        col_idx = -1
                        for row in reader:
                            if not row:
                                continue
                            if headers is None:
                                headers = [c.strip() for c in row]
                                for idx, h in enumerate(headers):
                                    if col_target in h.lower() or h.lower() in col_target:
                                        col_idx = idx
                                        matched_col_name = h
                                        break
                                if col_idx == -1:
                                    matched_col_name = headers[-1]
                                    col_idx = len(headers) - 1
                            else:
                                if 0 <= col_idx < len(row):
                                    val_str = row[col_idx].replace("$", "").replace("₹", "").replace(",", "").strip()
                                    try:
                                        values.append(float(val_str))
                                    except ValueError:
                                        pass

                if not values:
                    return PluginResult(
                        success=True,
                        data={
                            "file": os.path.basename(file_path),
                            "column": matched_col_name or col_target,
                            "operation": operation,
                            "count": 0,
                            "result": 0,
                            "formatted_result": "0",
                            "message": f"No numeric rows found in column '{matched_col_name or col_target}'."
                        }
                    )

                if operation in ("avg", "average"):
                    result = sum(values) / len(values)
                elif operation == "count":
                    result = len(values)
                elif operation == "max":
                    result = max(values)
                elif operation == "min":
                    result = min(values)
                else:
                    result = sum(values)

                formatted = f"{int(result):,}" if result.is_integer() else f"{result:,.2f}"
                return PluginResult(
                    success=True,
                    data={
                        "file": os.path.basename(file_path),
                        "path": file_path,
                        "column": matched_col_name or col_target,
                        "operation": operation,
                        "rows_count": len(values),
                        "result": result,
                        "formatted_result": formatted,
                        "message": f"Calculated {operation.upper()} of '{matched_col_name or col_target}': {formatted}."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Calculation failed: {e}")

        if action_name == "clean_sheet":
            file_path = str(parameters.get("path", "")).strip()
            operation = str(parameters.get("operation", "autofit")).strip().lower()

            if not file_path or not os.path.exists(file_path):
                return PluginResult(
                    success=True,
                    data={"operation": operation, "message": f"Operation '{operation}' ready."}
                )

            try:
                if file_path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.load_workbook(file_path)
                    ws = wb.active

                    # Autofit widths
                    if operation in ("autofit", "all"):
                        for col in ws.columns:
                            max_len = max(len(str(cell.value or "")) for cell in col)
                            col_letter = openpyxl.utils.get_column_letter(col[0].column)
                            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                    # Style headers
                    if operation in ("alignment", "borders", "all"):
                        thin_border = Border(
                            left=Side(style="thin", color="CCCCCC"),
                            right=Side(style="thin", color="CCCCCC"),
                            top=Side(style="thin", color="CCCCCC"),
                            bottom=Side(style="thin", color="CCCCCC"),
                        )
                        for cell in ws[1]:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
                            cell.alignment = Alignment(horizontal="center", vertical="center")

                    wb.save(file_path)

                return PluginResult(
                    success=True,
                    data={
                        "path": file_path,
                        "operation": operation,
                        "message": f"Successfully performed '{operation}' formatting on '{os.path.basename(file_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Clean sheet failed: {e}")

        if action_name == "create_chart":
            file_path = str(parameters.get("path", "")).strip()
            chart_type = str(parameters.get("chart_type", "bar")).strip().lower()

            if not file_path or not os.path.exists(file_path):
                return PluginResult(
                    success=True,
                    data={"chart_type": chart_type, "message": f"Prepared {chart_type.upper()} chart configuration."}
                )

            try:
                if file_path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.load_workbook(file_path)
                    ws = wb.active

                    if chart_type == "pie":
                        chart = PieChart()
                    elif chart_type == "line":
                        chart = LineChart()
                    else:
                        chart = BarChart()

                    chart.title = f"Data Chart ({chart_type.upper()})"
                    max_r = max(ws.max_row, 2)
                    data_ref = Reference(ws, min_col=2, min_row=1, max_row=max_r)
                    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=max_r)
                    chart.add_data(data_ref, titles_from_data=True)
                    chart.set_categories(cats_ref)
                    ws.add_chart(chart, "E2")
                    wb.save(file_path)

                return PluginResult(
                    success=True,
                    data={"chart_type": chart_type, "message": f"Created {chart_type.upper()} chart in '{os.path.basename(file_path)}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Chart creation failed: {e}")

        if action_name == "csv_to_excel":
            csv_path = str(parameters.get("csv_path", "")).strip()
            excel_path = str(parameters.get("excel_path", "")).strip()
            if not csv_path or not os.path.exists(csv_path):
                return PluginResult(success=False, error="Valid CSV path is required.")

            if not excel_path:
                excel_path = os.path.splitext(csv_path)[0] + ".xlsx"

            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Data"
                with open(csv_path, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.reader(f)
                    for row in reader:
                        ws.append(row)

                # Format header row
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                wb.save(excel_path)

                return PluginResult(
                    success=True,
                    data={
                        "csv_path": csv_path,
                        "excel_path": excel_path,
                        "rows": ws.max_row,
                        "message": f"Converted CSV to '{os.path.basename(excel_path)}' with {ws.max_row} rows."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"CSV to Excel conversion failed: {e}")

        if action_name == "excel_to_csv":
            excel_path = str(parameters.get("excel_path", "")).strip()
            csv_path = str(parameters.get("csv_path", "")).strip()
            sheet_name = str(parameters.get("sheet_name", "")).strip()

            if not excel_path or not os.path.exists(excel_path):
                return PluginResult(success=False, error="Valid Excel file path is required.")

            if not csv_path:
                csv_path = os.path.splitext(excel_path)[0] + ".csv"

            try:
                wb = openpyxl.load_workbook(excel_path, data_only=True)
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                row_count = 0
                with open(csv_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow([str(c) if c is not None else "" for c in row])
                        row_count += 1
                wb.close()

                return PluginResult(
                    success=True,
                    data={
                        "excel_path": excel_path,
                        "csv_path": csv_path,
                        "rows_exported": row_count,
                        "message": f"Exported {row_count} row(s) to '{os.path.basename(csv_path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Excel to CSV export failed: {e}")

        if action_name == "get_sheet_names":
            path = str(parameters.get("path", "")).strip()
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel file path is required.")

            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                sheets_info = []
                for sname in wb.sheetnames:
                    ws = wb[sname]
                    sheets_info.append({
                        "name": sname,
                        "max_row": ws.max_row,
                        "max_column": ws.max_column
                    })
                wb.close()

                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "sheets": sheets_info,
                        "count": len(sheets_info),
                        "message": f"Found {len(sheets_info)} worksheet(s) in '{os.path.basename(path)}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to get sheet names: {e}")

        if action_name == "merge_sheets":
            path = str(parameters.get("path", "")).strip()
            target_sheet = str(parameters.get("target_sheet", "Consolidated")).strip()
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel file path is required.")

            try:
                wb = openpyxl.load_workbook(path)
                if target_sheet in wb.sheetnames:
                    del wb[target_sheet]

                cons_ws = wb.create_sheet(title=target_sheet)
                total_merged_rows = 0
                header_written = False

                for sname in wb.sheetnames:
                    if sname == target_sheet:
                        continue
                    ws = wb[sname]
                    for idx, row in enumerate(ws.iter_rows(values_only=True)):
                        if idx == 0:
                            if not header_written:
                                cons_ws.append(list(row))
                                header_written = True
                        else:
                            cons_ws.append(list(row))
                            total_merged_rows += 1

                wb.save(path)
                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "target_sheet": target_sheet,
                        "merged_rows": total_merged_rows,
                        "message": f"Merged {total_merged_rows} data rows into sheet '{target_sheet}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to merge sheets: {e}")

        if action_name == "find_and_replace":
            path = str(parameters.get("path", "")).strip()
            find_text = str(parameters.get("find_text", ""))
            replace_text = str(parameters.get("replace_text", ""))
            sheet_name = str(parameters.get("sheet_name", "")).strip()

            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel file path is required.")

            try:
                wb = openpyxl.load_workbook(path)
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                replaced_count = 0

                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None and isinstance(cell.value, str):
                            if find_text in cell.value:
                                cell.value = cell.value.replace(find_text, replace_text)
                                replaced_count += 1

                wb.save(path)
                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "replaced_count": replaced_count,
                        "find_text": find_text,
                        "replace_text": replace_text,
                        "message": f"Replaced {replaced_count} occurrence(s) of '{find_text}' with '{replace_text}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Find and replace failed: {e}")

        if action_name == "add_formula":
            path = str(parameters.get("path", "")).strip()
            cell_coord = str(parameters.get("cell", "")).strip().upper()
            formula = str(parameters.get("formula", "")).strip()
            sheet_name = str(parameters.get("sheet_name", "")).strip()

            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel file path is required.")
            if not cell_coord or not formula:
                return PluginResult(success=False, error="'cell' and 'formula' parameters are required.")

            if not formula.startswith("="):
                formula = f"={formula}"

            try:
                wb = openpyxl.load_workbook(path)
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                ws[cell_coord] = formula
                wb.save(path)

                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "cell": cell_coord,
                        "formula": formula,
                        "message": f"Inserted formula '{formula}' into cell {cell_coord}."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to add formula: {e}")

        if action_name == "filter_rows":
            path = str(parameters.get("path", "")).strip()
            col_target = str(parameters.get("column_name", "")).strip().lower()
            condition = str(parameters.get("condition", "equals")).strip().lower()
            val_comp = str(parameters.get("value", "")).strip().lower()

            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel or CSV file path is required.")

            try:
                rows = []
                if path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.load_workbook(path, data_only=True)
                    ws = wb.active
                    for r in ws.iter_rows(values_only=True):
                        rows.append([str(c) if c is not None else "" for c in r])
                    wb.close()
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        rows = list(csv.reader(f))

                if not rows:
                    return PluginResult(success=True, data={"matches": [], "count": 0, "message": "Empty file."})

                headers = rows[0]
                col_idx = -1
                for idx, h in enumerate(headers):
                    if col_target in h.lower():
                        col_idx = idx
                        break
                if col_idx == -1:
                    col_idx = 0

                matches = []
                for r in rows[1:]:
                    if col_idx < len(r):
                        cell_val = str(r[col_idx]).strip().lower()
                        matched = False
                        if condition == "contains":
                            matched = val_comp in cell_val
                        elif condition in ("greater_than", "gt"):
                            try:
                                matched = float(cell_val) > float(val_comp)
                            except ValueError:
                                matched = False
                        elif condition in ("less_than", "lt"):
                            try:
                                matched = float(cell_val) < float(val_comp)
                            except ValueError:
                                matched = False
                        else:  # equals
                            matched = cell_val == val_comp

                        if matched:
                            matches.append(r)

                return PluginResult(
                    success=True,
                    data={
                        "headers": headers,
                        "matches": matches,
                        "count": len(matches),
                        "message": f"Found {len(matches)} matching row(s) for condition '{condition} {val_comp}'."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Filtering failed: {e}")

        if action_name == "sort_sheet":
            path = str(parameters.get("path", "")).strip()
            col_target = str(parameters.get("column_name", "")).strip().lower()
            ascending = bool(parameters.get("ascending", True))

            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel or CSV file path is required.")

            try:
                rows = []
                is_xlsx = path.lower().endswith((".xlsx", ".xlsm"))
                if is_xlsx:
                    wb = openpyxl.load_workbook(path)
                    ws = wb.active
                    for r in ws.iter_rows(values_only=True):
                        rows.append(list(r))
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        rows = list(csv.reader(f))

                if len(rows) <= 1:
                    return PluginResult(success=True, data={"message": "Insufficient data to sort."})

                headers = rows[0]
                col_idx = -1
                for idx, h in enumerate(headers):
                    if col_target in str(h).lower():
                        col_idx = idx
                        break
                if col_idx == -1:
                    col_idx = 0

                def sort_key(row):
                    if col_idx >= len(row) or row[col_idx] is None:
                        return (0, "")
                    val = row[col_idx]
                    try:
                        return (1, float(val))
                    except ValueError:
                        return (2, str(val).lower())

                data_rows = rows[1:]
                data_rows.sort(key=sort_key, reverse=not ascending)

                if is_xlsx:
                    ws.delete_rows(2, ws.max_row)
                    for r in data_rows:
                        ws.append(r)
                    wb.save(path)
                else:
                    with open(path, "w", newline="", encoding="utf-8") as f:
                        writer = csv.writer(f)
                        writer.writerow(headers)
                        writer.writerows(data_rows)

                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "sorted_column": str(headers[col_idx]),
                        "ascending": ascending,
                        "rows_sorted": len(data_rows),
                        "message": f"Sorted {len(data_rows)} rows by '{headers[col_idx]}' ({'ascending' if ascending else 'descending'})."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Sorting failed: {e}")

        if action_name == "export_sheet_to_json":
            path = str(parameters.get("path", "")).strip()
            out_json = str(parameters.get("output_json", "")).strip()
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel or CSV file path is required.")

            try:
                rows = []
                if path.lower().endswith((".xlsx", ".xlsm")):
                    wb = openpyxl.load_workbook(path, data_only=True)
                    ws = wb.active
                    for r in ws.iter_rows(values_only=True):
                        rows.append(list(r))
                    wb.close()
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        rows = list(csv.reader(f))

                if not rows:
                    return PluginResult(success=True, data={"items": [], "count": 0, "message": "Empty file."})

                headers = [str(h or f"Col_{i}") for i, h in enumerate(rows[0])]
                json_items = []
                for r in rows[1:]:
                    item = {}
                    for i, h in enumerate(headers):
                        item[h] = r[i] if i < len(r) else None
                    json_items.append(item)

                if out_json:
                    os.makedirs(os.path.dirname(os.path.abspath(out_json)), exist_ok=True)
                    with open(out_json, "w", encoding="utf-8") as f:
                        json.dump(json_items, f, indent=2)

                return PluginResult(
                    success=True,
                    data={
                        "path": path,
                        "output_json": out_json or None,
                        "count": len(json_items),
                        "items_preview": json_items[:5],
                        "message": f"Exported {len(json_items)} records to JSON."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"JSON export failed: {e}")

        if action_name == "get_workbook_summary":
            path = str(parameters.get("path", "")).strip()
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error="Valid Excel file path is required.")

            try:
                size_bytes = os.path.getsize(path)
                wb = openpyxl.load_workbook(path, data_only=False)
                sheets_meta = []
                total_cells = 0
                formula_count = 0

                for sname in wb.sheetnames:
                    ws = wb[sname]
                    sheet_formulas = 0
                    sheet_non_empty = 0
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.value is not None:
                                sheet_non_empty += 1
                                if isinstance(cell.value, str) and cell.value.startswith("="):
                                    sheet_formulas += 1

                    total_cells += (ws.max_row or 0) * (ws.max_column or 0)
                    formula_count += sheet_formulas
                    sheets_meta.append({
                        "name": sname,
                        "max_row": ws.max_row,
                        "max_col": ws.max_column,
                        "non_empty_cells": sheet_non_empty,
                        "formulas": sheet_formulas
                    })
                wb.close()

                return PluginResult(
                    success=True,
                    data={
                        "file": os.path.basename(path),
                        "size_bytes": size_bytes,
                        "sheets_count": len(sheets_meta),
                        "total_cells": total_cells,
                        "formula_count": formula_count,
                        "sheets": sheets_meta,
                        "message": f"Workbook '{os.path.basename(path)}': {len(sheets_meta)} sheet(s), {formula_count} formula(s), {size_bytes} bytes."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Workbook summary failed: {e}")

        return PluginResult(success=False, error=f"Unknown action '{action_name}' in Excel plugin.")

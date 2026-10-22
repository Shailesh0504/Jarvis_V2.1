# plugins/documents/plugin.py
"""Documents capability plugin implementation for reading, creating, editing, and analyzing documents."""

import datetime
import difflib
import logging
import os
import re
from typing import Any, Dict, List, Optional

from plugins.base import BasePlugin
from plugins.manifest import PluginManifest, PluginResult
from .commands import ACTIONS
from .permissions import PERMISSIONS_REQUIRED

logger = logging.getLogger("jarvis.plugins.documents")

try:
    import docx
except ImportError:
    docx = None


class DocumentsPlugin(BasePlugin):
    """Handles text, markdown, and Word (.docx) document reading, formatting, search, and diffing."""

    def get_manifest(self) -> PluginManifest:
        return PluginManifest(
            name="documents",
            version="2.0.0",
            description="Comprehensive text, markdown, and Word (.docx) document creator, editor, and analyzer.",
            author="Jarvis Team",
            permissions=PERMISSIONS_REQUIRED,
            actions=ACTIONS
        )

    def verify_state(self, action_name: str, parameters: Dict[str, Any], result_data: Any) -> Optional[Any]:
        from core.execution.models import VerificationResult, ExecutionState
        if action_name in ("create_document", "create_docx_document", "convert_document_format", "export_text_to_docx", "merge_text_documents"):
            path = (result_data or {}).get("filename") or (result_data or {}).get("file_path") or (result_data or {}).get("output_path") or (result_data or {}).get("destination_path")
            if path and os.path.isfile(path) and os.path.getsize(path) > 0:
                return VerificationResult(
                    verified=True,
                    state=ExecutionState.SUCCESS,
                    check_type="file_exists",
                    expected=path,
                    actual=f"{os.path.getsize(path)} bytes",
                    message="Document artifact verified on disk."
                )
            return VerificationResult(
                verified=False,
                state=ExecutionState.FAILED,
                check_type="file_exists",
                expected=path or "valid document",
                actual="Missing or empty",
                error="Expected document file not found."
            )
        return None

    def execute(self, action_name: str, parameters: Optional[Dict[str, Any]] = None) -> PluginResult:
        parameters = parameters or {}

        # 1. read_document
        if action_name == "read_document":
            path = parameters.get("file_path", "").strip()
            max_chars = int(parameters.get("max_chars", 2000))
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error=f"Document '{path}' nahi mila, Boss.")

            try:
                if path.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx library is not installed.")
                    doc = docx.Document(path)
                    text = "\n".join([p.text for p in doc.paragraphs if p.text])
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read(max_chars)

                preview = text[:300] + ("..." if len(text) > 300 else "")
                return PluginResult(
                    success=True,
                    data={"text": text, "preview": preview, "char_count": len(text), "message": f"Document content: {preview}"}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Error reading document: {e}")

        # 2. create_document
        elif action_name == "create_document":
            title = parameters.get("title", "Note").strip()
            content = parameters.get("content", "")
            fmt = parameters.get("format", "txt").lower().strip(".")
            safe_name = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).rstrip()
            filename = f"{safe_name or 'note'}.{fmt}"

            try:
                os.makedirs(os.path.dirname(os.path.abspath(filename)) or ".", exist_ok=True)
                with open(filename, "w", encoding="utf-8") as f:
                    f.write(content)
                return PluginResult(
                    success=True,
                    data={"filename": os.path.abspath(filename), "title": title, "bytes": len(content.encode('utf-8')), "message": f"Document '{filename}' save kar diya hai, Boss."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Error creating document: {e}")

        # 3. search_documents
        elif action_name == "search_documents":
            query = parameters.get("query", "").lower()
            dir_path = parameters.get("directory", ".")
            if not query:
                return PluginResult(success=False, error="query parameter is required.")

            matches = []
            try:
                for root, _, files in os.walk(dir_path):
                    for file in files:
                        if file.lower().endswith((".txt", ".md", ".docx", ".log", ".json")):
                            fp = os.path.join(root, file)
                            try:
                                if file.lower().endswith(".docx") and docx:
                                    doc = docx.Document(fp)
                                    text = " ".join(p.text for p in doc.paragraphs)
                                else:
                                    with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                                        text = f.read(100000)
                                if query in text.lower():
                                    matches.append(os.path.abspath(fp))
                            except Exception:
                                pass
                    if len(matches) >= 25:
                        break

                return PluginResult(
                    success=True,
                    data={"matches": matches, "count": len(matches), "message": f"Found {len(matches)} documents containing '{query}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Error searching documents: {e}")

        # 4. create_docx_document
        elif action_name == "create_docx_document":
            if not docx:
                return PluginResult(success=False, error="python-docx library is required.")

            file_path = parameters.get("file_path", "").strip()
            title = parameters.get("title", "").strip()
            paragraphs = parameters.get("paragraphs", [])
            bullet_points = parameters.get("bullet_points", [])

            if not file_path:
                return PluginResult(success=False, error="file_path is required.")
            if not file_path.lower().endswith(".docx"):
                file_path += ".docx"

            try:
                os.makedirs(os.path.dirname(os.path.abspath(file_path)) or ".", exist_ok=True)
                doc = docx.Document()
                if title:
                    doc.add_heading(title, level=0)

                for p in paragraphs:
                    if str(p).strip():
                        doc.add_paragraph(str(p))

                for bp in bullet_points:
                    if str(bp).strip():
                        doc.add_paragraph(str(bp), style='List Bullet')

                doc.save(file_path)
                return PluginResult(
                    success=True,
                    data={"file_path": os.path.abspath(file_path), "title": title, "paragraphs_count": len(paragraphs), "message": f"Created Word document: {file_path}"}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed to create docx: {e}")

        # 5. append_to_document
        elif action_name == "append_to_document":
            path = parameters.get("file_path", "").strip()
            content = parameters.get("content", "")
            if not path or not os.path.exists(path):
                return PluginResult(success=False, error=f"Target file '{path}' not found.")

            try:
                if path.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    doc = docx.Document(path)
                    doc.add_paragraph(content)
                    doc.save(path)
                else:
                    with open(path, "a", encoding="utf-8") as f:
                        f.write(f"\n{content}")

                return PluginResult(
                    success=True,
                    data={"file_path": os.path.abspath(path), "message": f"Appended content to '{path}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Failed appending to document: {e}")

        # 6. convert_document_format
        elif action_name == "convert_document_format":
            src = parameters.get("source_file", "").strip()
            target_fmt = parameters.get("target_format", "").strip().lower().strip(".")
            out_path = parameters.get("output_path", "").strip()

            if not src or not os.path.isfile(src):
                return PluginResult(success=False, error=f"Source file '{src}' not found.")

            if not out_path:
                base, _ = os.path.splitext(src)
                out_path = f"{base}.{target_fmt}"

            os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)

            try:
                # Read source text
                if src.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    doc = docx.Document(src)
                    text = "\n\n".join(p.text for p in doc.paragraphs if p.text)
                else:
                    with open(src, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read()

                # Write target
                if target_fmt == "docx":
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    new_doc = docx.Document()
                    for line in text.split("\n"):
                        if line.strip():
                            new_doc.add_paragraph(line)
                    new_doc.save(out_path)
                else:
                    with open(out_path, "w", encoding="utf-8") as f:
                        f.write(text)

                return PluginResult(
                    success=True,
                    data={"source_file": src, "output_path": os.path.abspath(out_path), "format": target_fmt, "message": f"Converted '{src}' to '{out_path}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Conversion error: {e}")

        # 7. get_document_statistics
        elif action_name == "get_document_statistics":
            path = parameters.get("file_path", "").strip()
            if not path or not os.path.isfile(path):
                return PluginResult(success=False, error=f"File '{path}' not found.")

            try:
                if path.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    doc = docx.Document(path)
                    text = " ".join(p.text for p in doc.paragraphs)
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read()

                words = len(re.findall(r'\b\w+\b', text))
                chars = len(text)
                chars_no_spaces = len(text.replace(" ", "").replace("\n", "").replace("\r", ""))
                lines = len(text.splitlines())
                sentences = len(re.findall(r'[.!?]+', text)) or 1
                reading_time_min = round(words / 200.0, 1)

                return PluginResult(
                    success=True,
                    data={
                        "file_path": os.path.abspath(path),
                        "words": words,
                        "characters": chars,
                        "characters_no_spaces": chars_no_spaces,
                        "lines": lines,
                        "sentences": sentences,
                        "reading_time_minutes": reading_time_min,
                        "message": f"Stats: {words} words, {lines} lines, ~{reading_time_min} min read."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Stats calculation error: {e}")

        # 8. search_and_replace_text
        elif action_name == "search_and_replace_text":
            path = parameters.get("file_path", "").strip()
            find_txt = parameters.get("find_text", "")
            replace_txt = parameters.get("replace_text", "")

            if not path or not os.path.isfile(path):
                return PluginResult(success=False, error=f"File '{path}' not found.")
            if not find_txt:
                return PluginResult(success=False, error="find_text is required.")

            try:
                if path.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    doc = docx.Document(path)
                    count = 0
                    for p in doc.paragraphs:
                        if find_txt in p.text:
                            count += p.text.count(find_txt)
                            p.text = p.text.replace(find_txt, replace_txt)
                    doc.save(path)
                else:
                    with open(path, "r", encoding="utf-8") as f:
                        content = f.read()
                    count = content.count(find_txt)
                    new_content = content.replace(find_txt, replace_txt)
                    with open(path, "w", encoding="utf-8") as f:
                        f.write(new_content)

                return PluginResult(
                    success=True,
                    data={"file_path": os.path.abspath(path), "replacements_count": count, "message": f"Replaced {count} occurrence(s) in '{path}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Search and replace error: {e}")

        # 9. compare_documents
        elif action_name == "compare_documents":
            fa = parameters.get("file_a", "").strip()
            fb = parameters.get("file_b", "").strip()

            if not fa or not os.path.isfile(fa):
                return PluginResult(success=False, error=f"File A '{fa}' not found.")
            if not fb or not os.path.isfile(fb):
                return PluginResult(success=False, error=f"File B '{fb}' not found.")

            try:
                with open(fa, "r", encoding="utf-8", errors="ignore") as f1:
                    lines_a = f1.readlines()
                with open(fb, "r", encoding="utf-8", errors="ignore") as f2:
                    lines_b = f2.readlines()

                diff = list(difflib.unified_diff(lines_a, lines_b, fromfile=os.path.basename(fa), tofile=os.path.basename(fb)))
                added = sum(1 for line in diff if line.startswith("+") and not line.startswith("+++"))
                removed = sum(1 for line in diff if line.startswith("-") and not line.startswith("---"))

                return PluginResult(
                    success=True,
                    data={
                        "file_a": fa,
                        "file_b": fb,
                        "identical": len(diff) == 0,
                        "added_lines": added,
                        "removed_lines": removed,
                        "diff_sample": diff[:20],
                        "message": "Files are identical." if len(diff) == 0 else f"Diff: +{added} lines, -{removed} lines."
                    }
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Comparison error: {e}")

        # 10. extract_document_sections
        elif action_name == "extract_document_sections":
            path = parameters.get("file_path", "").strip()
            if not path or not os.path.isfile(path):
                return PluginResult(success=False, error=f"File '{path}' not found.")

            sections = []
            try:
                if path.lower().endswith(".docx"):
                    if not docx:
                        return PluginResult(success=False, error="python-docx is required.")
                    doc = docx.Document(path)
                    for p in doc.paragraphs:
                        if p.style and "Heading" in p.style.name:
                            sections.append({"title": p.text, "level": p.style.name})
                else:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        for line in f:
                            m = re.match(r'^(#{1,6})\s+(.*)', line.strip())
                            if m:
                                sections.append({"level": len(m.group(1)), "title": m.group(2)})

                return PluginResult(
                    success=True,
                    data={"file_path": os.path.abspath(path), "section_count": len(sections), "sections": sections, "message": f"Extracted {len(sections)} section heading(s)."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Section extraction error: {e}")

        # 11. export_text_to_docx
        elif action_name == "export_text_to_docx":
            if not docx:
                return PluginResult(success=False, error="python-docx is required.")

            text = parameters.get("text", "")
            dest = parameters.get("destination_path", "").strip()
            title = parameters.get("title", "Exported Document").strip()

            if not dest:
                return PluginResult(success=False, error="destination_path is required.")
            if not dest.lower().endswith(".docx"):
                dest += ".docx"

            try:
                os.makedirs(os.path.dirname(os.path.abspath(dest)) or ".", exist_ok=True)
                doc = docx.Document()
                if title:
                    doc.add_heading(title, level=0)
                for paragraph in text.split("\n\n"):
                    if paragraph.strip():
                        doc.add_paragraph(paragraph.strip())
                doc.save(dest)

                return PluginResult(
                    success=True,
                    data={"destination_path": os.path.abspath(dest), "title": title, "message": f"Exported text to '{dest}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Docx export error: {e}")

        # 12. merge_text_documents
        elif action_name == "merge_text_documents":
            file_paths = parameters.get("file_paths", [])
            out_path = parameters.get("output_path", "").strip()

            if not file_paths:
                return PluginResult(success=False, error="file_paths list is required.")
            if not out_path:
                return PluginResult(success=False, error="output_path is required.")

            os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
            merged_content = []

            try:
                for fp in file_paths:
                    if os.path.isfile(fp):
                        merged_content.append(f"\n--- {os.path.basename(fp)} ---\n")
                        with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                            merged_content.append(f.read())

                with open(out_path, "w", encoding="utf-8") as f:
                    f.write("\n".join(merged_content))

                return PluginResult(
                    success=True,
                    data={"output_path": os.path.abspath(out_path), "merged_count": len(file_paths), "message": f"Merged {len(file_paths)} files into '{out_path}'."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Merge error: {e}")

        # 13. inspect_document_metadata
        elif action_name == "inspect_document_metadata":
            path = parameters.get("file_path", "").strip()
            if not path or not os.path.isfile(path):
                return PluginResult(success=False, error=f"File '{path}' not found.")

            try:
                st = os.stat(path)
                mtime = datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
                size_kb = round(st.st_size / 1024.0, 2)
                meta = {
                    "file_name": os.path.basename(path),
                    "path": os.path.abspath(path),
                    "size_kb": size_kb,
                    "modified": mtime,
                    "extension": os.path.splitext(path)[1]
                }

                if path.lower().endswith(".docx") and docx:
                    doc = docx.Document(path)
                    core_props = doc.core_properties
                    meta["docx_author"] = core_props.author or "Unknown"
                    meta["docx_created"] = str(core_props.created or "")
                    meta["paragraphs_count"] = len(doc.paragraphs)

                return PluginResult(
                    success=True,
                    data={"metadata": meta, "message": f"Document '{os.path.basename(path)}': {size_kb} KB, modified {mtime}."}
                )
            except Exception as e:
                return PluginResult(success=False, error=f"Metadata error: {e}")

        return PluginResult(success=False, error=f"Unknown action '{action_name}' in Documents plugin.")
